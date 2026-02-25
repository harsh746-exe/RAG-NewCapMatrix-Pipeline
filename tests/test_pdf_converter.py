import os
import sys
import tempfile
import unittest
from pathlib import Path


# Ensure `src/` is importable when running `python -m unittest`
PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))


from core.pdf_converter import PDFToTXTConverter  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402


class TestPDFConverter(unittest.TestCase):
    def test_discovers_pdfs_and_docx_recursively(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir) / "02_Detailed Project Descriptions"
            contract_a = root / "Contract A"
            contract_b = root / "Contract B" / "sub"
            contract_a.mkdir(parents=True)
            contract_b.mkdir(parents=True)

            (contract_a / "a1.pdf").write_bytes(b"%PDF-1.4\n%fake\n")
            (contract_b / "b1.pdf").write_bytes(b"%PDF-1.4\n%fake\n")
            (contract_a / "a1.docx").write_bytes(b"")  # file exists; content not needed for discovery
            (contract_b / "b1.docx").write_bytes(b"")

            converter = PDFToTXTConverter(pdf_dir=str(root), txt_dir=str(Path(tmpdir) / "documents_txt"))
            pdfs = converter.get_pdf_files()
            docxs = converter.get_docx_files()

            self.assertEqual({p.name for p in pdfs}, {"a1.pdf", "b1.pdf"})
            self.assertEqual({p.name for p in docxs}, {"a1.docx", "b1.docx"})

    def test_converts_docx_to_txt_with_header(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir) / "02_Detailed Project Descriptions" / "Contract X"
            root.mkdir(parents=True)
            txt_out = Path(tmpdir) / "documents_txt"

            docx_path = root / "sample.docx"
            doc = DocxDocument()
            doc.add_paragraph("Hello from DOCX.")
            doc.add_paragraph("Second paragraph.")
            doc.save(docx_path)

            converter = PDFToTXTConverter(pdf_dir=str(Path(tmpdir) / "02_Detailed Project Descriptions"), txt_dir=str(txt_out))
            result = converter.convert_docx_to_txt(docx_path)
            self.assertTrue(result["success"])

            txt_path = txt_out / "sample.txt"
            self.assertTrue(txt_path.exists())
            content = txt_path.read_text(encoding="utf-8")
            self.assertIn("# Original DOCX:", content)
            self.assertIn("Hello from DOCX.", content)


if __name__ == "__main__":
    unittest.main()

